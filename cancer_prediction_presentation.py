from pptx import Presentation
from pptx.util import Inches, Pt

# Create a presentation
prs = Presentation()

# Slide Titles and Content
slides = [
    ("Cancer Disease Prediction Using ML, FastAPI, Streamlit, Docker & CI/CD",
     "Author: Mary Sophiya\nDate: 2025-12-30"),
    ("Problem Statement",
     "Early cancer detection is critical.\nObjective: Build an ML system for Malignant vs Benign prediction."),
    ("Dataset & Features",
     "Dataset: scikit-learn breast_cancer dataset\nFeatures: 30 numeric tumor characteristics\nTarget: 0=Malignant, 1=Benign"),
    ("ML Pipeline",
     "Data Loading → Preprocessing → Model Training → Model Saving\nModel: RandomForestClassifier\nScaler: StandardScaler"),
    ("Prediction Module",
     "Lazy model loading\nScaling input\nPrediction: 'Malignant' / 'Benign'\nPrevents import-time errors during testing"),
    ("FastAPI Backend",
     "Endpoint: /predict\nReceives JSON input → returns prediction\nEnables programmatic access"),
    ("Streamlit UI",
     "Interactive UI for user input\nSends request to API → displays prediction"),
    ("Version Control",
     "Git + GitHub\nBranching Strategy: GitHub Flow\nFeature branches → PR → Merge to main\n.gitignore used to exclude unnecessary files"),
    ("Testing Approach",
     "Unit Tests: ML pipeline functions\nIntegration Tests: Full workflow train → predict\nCI/CD Integration: Tests run automatically"),
    ("CI/CD Pipeline",
     "GitHub Actions Workflow:\n1. Checkout repo\n2. Install dependencies\n3. Train model\n4. Lint & run tests\n5. Build Docker images (API & UI)\n6. Push to Docker Hub"),
    ("Docker & Deployment",
     "Docker images ensure consistent environment\nDocker Compose runs API + UI together\nPublic Docker Hub images can be pulled anywhere"),
    ("Best Practices Applied",
     "Version control & branching\nLazy loading of ML models\nAutomated testing with pytest\nCI/CD with Docker & GitHub Actions\nSecure handling of secrets"),
    ("Results & Screenshots",
     "Model Accuracy: ~95%\nAPI response example\nStreamlit UI screenshot"),
    ("Conclusion & Future Work",
     "ML-based cancer prediction successful\nStreamlit + FastAPI + Docker + CI/CD integrated\nFuture: Ensemble models, cloud deployment, monitoring"),
    ("References",
     "scikit-learn: https://scikit-learn.org\nFastAPI: https://fastapi.tiangolo.com\nStreamlit: https://streamlit.io\nDocker: https://www.docker.com")
]

# Add slides to presentation
for title, content in slides:
    slide_layout = prs.slide_layouts[1]  # Title + Content
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = title
    slide.placeholders[1].text = content

# Save presentation
prs.save('Cancer_Prediction_Project_Presentation.pptx')

print("Presentation created: Cancer_Prediction_Project_Presentation.pptx")
